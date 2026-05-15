import json
import re
import zipfile
from datetime import datetime
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DB_PATH = ROOT / "data" / "database.json"

AUDITS = {
    "ATS": Path(r"C:\Users\micro\OneDrive\Desktop\Monthly Audit from Claude\Amazon Shipping\DERMATOUCH_ATS_Audit_Apr2026.xlsx"),
    "Delhivery": Path(r"C:\Users\micro\OneDrive\Desktop\Monthly Audit from Claude\Delivery partner Audit\CWPL_Delhivery_MAY2026_Audit_v2.xlsx"),
    "DTDC": Path(r"C:\Users\micro\OneDrive\Desktop\Monthly Audit from Claude\DTDC Audit\DTDC_Billing_Audit_April2026.xlsx"),
    "GoSwift": Path(r"C:\Users\micro\OneDrive\Desktop\Monthly Audit from Claude\Go Swift Audit\GoSwift_Billing_Audit_Report_May2026.xlsx"),
    "Shadowfax": Path(r"C:\Users\micro\OneDrive\Desktop\Monthly Audit from Claude\Shadofax Audit\Dermatouch_MAY2026_Audit_Report.xlsx"),
    "XpressBees": Path(r"C:\Users\micro\OneDrive\Desktop\Monthly Audit from Claude\Xpressbees\Xpressbees_MAY2026_Audit_v2.xlsx"),
}


def money(value):
    text = str(value or "")
    text = text.replace(",", "")
    match = re.search(r"-?\d+(?:\.\d+)?", text)
    return round(float(match.group()), 2) if match else 0.0


def intish(value):
    text = str(value or "").replace(",", "")
    match = re.search(r"\d+", text)
    return int(match.group()) if match else 0


def open_rows(path, sheet_name):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet_name]
    return [[cell for cell in row] for row in ws.iter_rows(values_only=True)]


def xml_workbook(path):
    ns = {
        "m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
        "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    }
    with zipfile.ZipFile(path) as zf:
        names = zf.namelist()
        shared = []
        if "xl/sharedStrings.xml" in names:
            root = ET.fromstring(zf.read("xl/sharedStrings.xml"))
            for si in root.findall("m:si", ns):
                shared.append("".join(t.text or "" for t in si.findall(".//m:t", ns)))

        workbook = ET.fromstring(zf.read("xl/workbook.xml"))
        rels = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
        rid_to_target = {rel.attrib["Id"]: rel.attrib["Target"] for rel in rels}
        sheets = {}

        for sheet in workbook.findall("m:sheets/m:sheet", ns):
            title = sheet.attrib["name"]
            rid = sheet.attrib["{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"]
            target = rid_to_target[rid].lstrip("/")
            if not target.startswith("xl/"):
                target = "xl/" + target
            root = ET.fromstring(zf.read(target))
            rows = []
            for row in root.findall(".//m:sheetData/m:row", ns):
                values = []
                for cell in row.findall("m:c", ns):
                    value_node = cell.find("m:v", ns)
                    value = "" if value_node is None else value_node.text or ""
                    if cell.attrib.get("t") == "s" and value.isdigit() and int(value) < len(shared):
                        value = shared[int(value)]
                    values.append(value)
                rows.append(values)
            sheets[title] = rows
        return sheets


def audit_ats():
    rows = open_rows(AUDITS["ATS"], "Executive Summary")
    claim_ex = claim = 0
    for row in rows:
        if str(row[0] or "").startswith("TOTAL CONFIRMED OVERCHARGE"):
            claim_ex = money(row[4])
            claim = money(row[5])
    return {
        "carrier": "Amazon ATS",
        "period": "April 2026",
        "shipmentsAudited": 21555,
        "errorCount": 54,
        "claimAmountExGst": claim_ex,
        "claimAmount": claim,
        "trusted": True,
        "confidence": "confirmed",
        "sourceFile": str(AUDITS["ATS"]),
    }


def audit_delhivery():
    sheets = xml_workbook(AUDITS["Delhivery"])
    summary = sheets["Executive Summary"]
    row = summary[4]
    return {
        "carrier": "Delhivery",
        "period": "May 2026",
        "shipmentsAudited": intish(row[0]),
        "errorCount": intish(row[1]),
        "claimAmountExGst": money(row[3]),
        "claimAmount": money(row[4]),
        "trusted": True,
        "confidence": "confirmed",
        "sourceFile": str(AUDITS["Delhivery"]),
    }


def audit_xpressbees():
    sheets = xml_workbook(AUDITS["XpressBees"])
    row = sheets["Executive Summary"][4]
    return {
        "carrier": "XpressBees",
        "period": "May 2026",
        "shipmentsAudited": intish(row[0]),
        "errorCount": intish(row[1]),
        "claimAmountExGst": money(row[3]),
        "claimAmount": money(row[4]),
        "trusted": True,
        "confidence": "confirmed",
        "sourceFile": str(AUDITS["XpressBees"]),
    }


def audit_goswift():
    rows = open_rows(AUDITS["GoSwift"], "Dispute Summary")
    shipments = errors = 0
    claim = 0
    for row in rows:
        text = " ".join(str(v or "") for v in row)
        if "Total Shipments Audited" in text:
            shipments = intish(text)
        if "Shipments with Billing Errors" in text:
            errors = intish(text)
        if "TOTAL REFUND CLAIM" in text:
            claim = money(text)
    return {
        "carrier": "GoSwift",
        "period": "May 2026",
        "shipmentsAudited": shipments,
        "errorCount": errors,
        "claimAmountExGst": claim,
        "claimAmount": claim,
        "trusted": True,
        "confidence": "confirmed",
        "sourceFile": str(AUDITS["GoSwift"]),
    }


def audit_shadowfax():
    rows = open_rows(AUDITS["Shadowfax"], "Dashboard")
    row = rows[5]
    return {
        "carrier": "Shadowfax",
        "period": "May 2026",
        "shipmentsAudited": intish(row[0]),
        "errorCount": intish(row[2]),
        "claimAmountExGst": money(row[8]),
        "claimAmount": money(row[8]),
        "trusted": True,
        "confidence": "confirmed",
        "sourceFile": str(AUDITS["Shadowfax"]),
    }


def audit_dtdc():
    rows = open_rows(AUDITS["DTDC"], "Dispute Summary")
    shipments = errors = 0
    claim = 0
    for row in rows:
        label = str(row[0] or "")
        if label == "Total Shipments Audited":
            shipments = intish(row[1])
        if label == "Shipments with Billing Errors":
            errors = intish(row[1])
        if "TOTAL REFUND CLAIM" in label:
            claim = money(row[1])
    return {
        "carrier": "DTDC",
        "period": "April 2026",
        "shipmentsAudited": shipments,
        "errorCount": errors,
        "claimAmountExGst": claim,
        "claimAmount": claim,
        "trusted": False,
        "confidence": "excluded_by_user_review",
        "sourceFile": str(AUDITS["DTDC"]),
    }


audits = [audit_ats(), audit_delhivery(), audit_xpressbees(), audit_goswift(), audit_shadowfax(), audit_dtdc()]

with DB_PATH.open("r", encoding="utf-8") as handle:
    data = json.load(handle)

data["referenceAudits"] = audits
data.setdefault("auditLogs", []).append({
    "id": f"aud_reference_{datetime.utcnow().isoformat()}",
    "actorId": "demo_user",
    "action": "reference_audits_applied",
    "entityType": "referenceAudit",
    "entityId": None,
    "beforeState": None,
    "afterState": {
        "trustedCarriers": [audit["carrier"] for audit in audits if audit["trusted"]],
        "excludedCarriers": [audit["carrier"] for audit in audits if not audit["trusted"]],
        "confirmedClaimAmount": round(sum(audit["claimAmount"] for audit in audits if audit["trusted"]), 2),
    },
    "createdAt": datetime.utcnow().isoformat() + "Z",
})

with DB_PATH.open("w", encoding="utf-8") as handle:
    json.dump(data, handle, indent=2)

print(json.dumps({
    "applied": audits,
    "confirmedClaimAmount": round(sum(audit["claimAmount"] for audit in audits if audit["trusted"]), 2),
}, indent=2))
