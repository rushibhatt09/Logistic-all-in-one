import json
import sys
from openpyxl import load_workbook


def cell_value(value):
    if value is None:
        return ""
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return value


if len(sys.argv) < 2:
    raise SystemExit("Missing Excel file path")

path = sys.argv[1]
workbook = load_workbook(path, read_only=True, data_only=True)
keywords = {
    "awb", "waybill", "tracking", "order", "status", "carrier", "courier",
    "zone", "weight", "amount", "charge", "total", "pincode", "shipment"
}


def header_score(values):
    text = " ".join(str(value).strip().lower() for value in values if value is not None)
    return sum(1 for keyword in keywords if keyword in text)


best = None
for sheet in workbook.worksheets:
    for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        values = [cell_value(value) for value in row]
        if not any(str(value).strip() for value in values):
            continue
        score = header_score(values)
        if best is None or score > best["score"]:
            best = {"sheet": sheet, "row_number": row_number, "headers": values, "score": score}
        if row_number > 30:
            break

if not best:
    print("[]")
    raise SystemExit(0)

sheet = best["sheet"]
headers = [str(value).strip() if value is not None else "" for value in best["headers"]]
rows = sheet.iter_rows(min_row=best["row_number"] + 1, values_only=True)

if not headers:
    print("[]")
    raise SystemExit(0)

records = []
for row in rows:
    values = [cell_value(value) for value in row]
    if not any(str(value).strip() for value in values):
        continue
    record = {}
    for index, header in enumerate(headers):
        if not header:
            continue
        record[header] = values[index] if index < len(values) else ""
    records.append(record)

print(json.dumps(records, ensure_ascii=False))
